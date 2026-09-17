"""Unit tests for TSO Portaria overlay adapters (TAG / TBG / NTS)."""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "flows"))

from tso import base  # noqa: E402
from tso import nts as nts_mod  # noqa: E402
from tso import tag as tag_mod  # noqa: E402
from tso import tbg as tbg_mod  # noqa: E402


def test_prefer_newest_prog_real_files(tmp_path: Path):
    older = tmp_path / "Prog-Real_GASENE_01-2017_a_07-2026.xlsx"
    newer = tmp_path / "Prog-Real_GASENE_01-2017_a_08-2026.xlsx"
    other = tmp_path / "Prog-Real_MALHAS-NE_01-2017_a_07-2026.xlsx"
    for p in (older, newer, other):
        p.write_bytes(b"PK")  # content unused
    kept = {p.name for p in base.prefer_newest_prog_real_files([older, newer, other])}
    assert "Prog-Real_GASENE_01-2017_a_08-2026.xlsx" in kept
    assert "Prog-Real_GASENE_01-2017_a_07-2026.xlsx" not in kept
    assert "Prog-Real_MALHAS-NE_01-2017_a_07-2026.xlsx" in kept


def test_tag_build_prefers_newest_coverage(tmp_path: Path, monkeypatch):
    """Regression: cached July + August workbooks must not both melt."""
    calls: list[str] = []

    def fake_parse(path, **kwargs):
        calls.append(path.name)
        return pd.DataFrame(
            [{
                "date": pd.Timestamp("2026-08-01"),
                "point_code": "x",
                "point_name": "P",
                "point_type": base.POINT_TYPE_DELIVERY,
                "pipeline_code": "tag:gasene",
                "pipeline_name": "GASENE",
                "municipality": None,
                "uf": "RJ",
                "tso": "TAG",
                "variable": base.VAR_ACTUAL,
                "value": 1.0,
                "source": "tag",
            }]
        )

    monkeypatch.setattr(base, "parse_wide_prog_real_workbook", fake_parse)
    (tmp_path / "Prog-Real_GASENE_01-2017_a_07-2026.xlsx").write_bytes(b"PK")
    (tmp_path / "Prog-Real_GASENE_01-2017_a_08-2026.xlsx").write_bytes(b"PK")
    df = tag_mod.build(tmp_path)
    assert len(df) == 1
    assert calls == ["Prog-Real_GASENE_01-2017_a_08-2026.xlsx"]


def _write_tbg_entregues(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Prog x Real Saídas"
    ws["B3"] = pd.Timestamp("2026-08-01")
    ws["F3"] = "Volumes Entregues"
    ws["C7"] = "Corumbá"
    ws["E7"] = "Campo Grande"
    ws["C8"] = 19581
    ws["E8"] = 26983
    ws["C9"] = "Programado"
    ws["D9"] = "Realizado"
    ws["E9"] = "Programado"
    ws["F9"] = "Realizado"
    ws["B10"] = pd.Timestamp("2026-08-01")
    ws["C10"] = 10
    ws["D10"] = 11
    ws["E10"] = 20
    ws["F10"] = 21
    ws["B11"] = pd.Timestamp("2026-08-02")
    ws["C11"] = 12
    ws["D11"] = 13
    ws["E11"] = 22
    ws["F11"] = 23
    wb.save(path)


def test_tbg_parse_paired_entregues(tmp_path: Path):
    path = tmp_path / "Volumes_Entregues_Agosto.xlsx"
    _write_tbg_entregues(path)
    df = tbg_mod.build(tmp_path)
    assert not df.empty
    assert set(df["variable"]) == {base.VAR_ACTUAL, base.VAR_SCHEDULED}
    assert df["date"].max() == pd.Timestamp("2026-08-02")
    # Values already thousand m³ — not divided again.
    corumba_actual = df[
        (df["point_name"] == "Corumbá") & (df["variable"] == base.VAR_ACTUAL) & (df["date"] == "2026-08-01")
    ]
    assert abs(float(corumba_actual.iloc[0]["value"]) - 11.0) < 1e-9
    assert "19581" in set(df["point_code"].astype(str))


def _write_nts_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Programado (AGO)"
    ws["F3"] = "RELATÓRIO DOS VOLUMES PROGRAMADOS - 2026"
    ws["D11"] = "PTE BARBACENA"
    ws["E11"] = "PTE BETIM II"
    ws["D12"] = "(mil m³/dia)"
    ws["E12"] = "(mil m³/dia)"
    ws["C13"] = pd.Timestamp("2026-08-01")
    ws["D13"] = 1.5
    ws["E13"] = 2.5
    ws["C14"] = pd.Timestamp("2026-08-02")
    ws["D14"] = 3.0
    ws["E14"] = 4.0
    ws2 = wb.create_sheet("Realizado (AGO)")
    ws2["D11"] = "PTE BARBACENA"
    ws2["E11"] = "PTE BETIM II"
    ws2["D12"] = "(mil m³/dia)"
    ws2["E12"] = "(mil m³/dia)"
    ws2["C13"] = pd.Timestamp("2026-08-01")
    ws2["D13"] = 9.0
    ws2["E13"] = 8.0
    wb.save(path)


def test_nts_parse_monthly_sheets(tmp_path: Path):
    path = tmp_path / "nts_Quantidade_Programada_2026.xlsx"
    _write_nts_workbook(path)
    # Also need a realizado file name cue — sheets carry the variable.
    df = nts_mod.build(tmp_path)
    assert not df.empty
    assert df["date"].max() == pd.Timestamp("2026-08-02")
    assert base.VAR_SCHEDULED in set(df["variable"])
    assert base.VAR_ACTUAL in set(df["variable"])
    sched = df[(df["point_name"] == "PTE BARBACENA") & (df["variable"] == base.VAR_SCHEDULED)]
    assert abs(float(sched.iloc[0]["value"]) - 1.5) < 1e-9


def test_tbg_filename_from_liferay_url():
    url = (
        "https://www.tbg.com.br/documents/20124/657762/"
        "ANP+Agosto+2026.zip/e1baa0a4-c095-1754-06f1-de4514fe14e3?t=1"
    )
    assert tbg_mod._filename_from_url(url, "fallback.bin") == "ANP_Agosto_2026.zip"

    text = (ROOT / ".github/workflows/flows.yml").read_text(encoding="utf-8")
    assert "cron: \"15 12 * * 1,4\"" in text
    assert "workflow_dispatch" in text
    assert "raw/tso" in text
